#!/usr/bin/env python3
"""Generate FunctionalTestCase.xlsx from FunctionalTestCase.csv."""

import csv
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
CSV_PATH = ROOT / "FunctionalTestCase.csv"
XLSX_PATH = ROOT / "FunctionalTestCase.xlsx"

HEADERS = [
    "Test Case ID",
    "Module",
    "Title",
    "Requirement ID",
    "Priority",
    "Suite Type",
    "Preconditions",
    "Test Data",
    "Test Steps",
    "Expected Result",
    "Status",
    "Executed By",
    "Execution Date",
    "Comments",
]

HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")


def load_rows():
    with CSV_PATH.open(newline="", encoding="utf-8") as handle:
        reader = csv.DictReader(handle)
        return list(reader)


def autosize_columns(sheet, max_width=60):
    for column_cells in sheet.columns:
        letter = get_column_letter(column_cells[0].column)
        length = max(len(str(cell.value or "")) for cell in column_cells)
        sheet.column_dimensions[letter].width = min(max(length + 2, 12), max_width)


def write_functional_test_cases(sheet, rows):
    sheet.title = "FunctionalTestCase"
    sheet.append(HEADERS)

    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in rows:
        sheet.append(
            [
                row.get("Test Case ID", ""),
                row.get("Module", ""),
                row.get("Title", ""),
                row.get("Requirement ID", ""),
                row.get("Priority", ""),
                row.get("Suite Type", ""),
                row.get("Preconditions", ""),
                row.get("Test Data", ""),
                row.get("Test Steps", ""),
                row.get("Expected Result", ""),
                "",
                "",
                "",
                "",
            ]
        )

    for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row):
        for cell in row:
            cell.alignment = WRAP

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(sheet.max_column)}{sheet.max_row}"
    autosize_columns(sheet)


def write_summary(sheet, rows):
    sheet.title = "Summary"
    sheet.append(["FunctionalTestCase — Summary"])
    sheet["A1"].font = Font(bold=True, size=14)

    module_counts = Counter(row["Module"] for row in rows)
    suite_counts = Counter(row["Suite Type"] for row in rows)
    priority_counts = Counter(row["Priority"] for row in rows)

    sheet.append([])
    sheet.append(["Total Test Cases", len(rows)])
    sheet.append([])
    sheet.append(["By Module", "Count"])
    for module, count in sorted(module_counts.items()):
        sheet.append([module, count])

    sheet.append([])
    sheet.append(["By Suite Type", "Count"])
    for suite, count in sorted(suite_counts.items()):
        sheet.append([suite, count])

    sheet.append([])
    sheet.append(["By Priority", "Count"])
    for priority in ("P1", "P2", "P3"):
        sheet.append([priority, priority_counts.get(priority, 0)])

    sheet.column_dimensions["A"].width = 28
    sheet.column_dimensions["B"].width = 12


def write_requirement_coverage(sheet, rows):
    sheet.title = "Requirement Coverage"
    sheet.append(["Requirement ID", "Test Case ID(s)"])
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    coverage = {}
    for row in rows:
        req_ids = [part.strip() for part in row["Requirement ID"].split(",")]
        test_id = row["Test Case ID"]
        for req_id in req_ids:
            coverage.setdefault(req_id, []).append(test_id)

    for req_id in sorted(coverage):
        sheet.append([req_id, ", ".join(coverage[req_id])])

    for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row):
        for cell in row:
            cell.alignment = WRAP

    sheet.freeze_panes = "A2"
    autosize_columns(sheet)


def main():
    rows = load_rows()
    workbook = Workbook()
    write_functional_test_cases(workbook.active, rows)
    write_summary(workbook.create_sheet(), rows)
    write_requirement_coverage(workbook.create_sheet(), rows)
    workbook.save(XLSX_PATH)
    print(f"Created {XLSX_PATH} ({len(rows)} test cases)")


if __name__ == "__main__":
    main()
